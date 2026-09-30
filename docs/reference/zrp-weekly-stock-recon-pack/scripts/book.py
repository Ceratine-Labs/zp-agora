# executed inside run_weekly.py - builds the master book from the run's data
import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from matrix import MATRIX, SHIFT_TESTS, RULES
from playbook import BUCKET, BUCKET_ORDER, PLAY

# ------------------------------------------------------------------ data
g = g.reset_index(drop=True)
SITES = sorted(g.BranchName.unique())
class _Short(dict):
    def __missing__(self, k): return k
SHORT = _Short({'Baobab Convenience Centre': 'Baobab CC', 'Baobab Inn': 'Baobab Inn', 'Captains Convenience Centre': 'Captains', 'Engen Bethlehem One Plus': 'Engen Bethlehem',
         'Engen Manguzi': 'Engen Manguzi', 'Ngwenya Service Station': 'Ngwenya', 'Nyala One Stop': 'Nyala', 'Theku Plaza (Cygnisol)': 'Theku Plaza', 'Total Hluhluwe': 'Total Hluhluwe'})
def _lab(a, b):
    return f'{a.day}-{b.day} {b:%B %Y}' if a.month == b.month else f'{a.day} {a:%B} - {b.day} {b:%B %Y}'
WL = _lab(WF, WT); PL = _lab(PERIOD_FROM, PERIOD_TO); WLs = f'{WF.day}-{WT.day} {WT:%b}' if WF.month == WT.month else f'{WF:%d %b}-{WT:%d %b}'
PLs = f'{PERIOD_FROM.day}-{PERIOD_TO.day} {PERIOD_TO:%b}' if PERIOD_FROM.month == PERIOD_TO.month else f'{PERIOD_FROM:%d %b}-{PERIOD_TO:%d %b}'
PTAB = ('Period Items ' + PLs)[:31]

# opening vs closing register, 1-21 Sep, corrected data
sh = r.sort_values(['BranchName', 'AreaNo', 'StockItemNo', 'Date', 'ShiftNo']).copy()
gg = sh.groupby(['BranchName', 'AreaNo', 'StockItemNo'])
sh['PrevCloseStaff'] = gg.QtyClose_Original.shift(1); sh['PrevCloseBal'] = gg.QtyClose_Amended.shift(1)
sh = sh[sh.PrevCloseStaff.notna()]
sh['Break'] = sh.QtyOpen_Original - sh.PrevCloseStaff
cb = sh[sh.Break.abs() > 1e-6].copy()
cb['Repair'] = np.select([(cb.QtyOpen_Amended - cb.PrevCloseBal).abs() > 1e-6,
    ((cb.PrevCloseBal - cb.PrevCloseStaff).abs() > 1e-6) & ((cb.QtyOpen_Amended - cb.QtyOpen_Original).abs() <= 1e-6),
    ((cb.QtyOpen_Amended - cb.QtyOpen_Original).abs() > 1e-6) & ((cb.PrevCloseBal - cb.PrevCloseStaff).abs() <= 1e-6)],
    ['Still open after balancing', 'WRONG - previous closing changed', 'Correct - opening amended'], 'Both changed')
cb['Direction'] = np.where(cb.Break < 0, 'Opening BELOW previous closing', 'Opening ABOVE previous closing')
cb['Value'] = cb.Break * cb.RandPU
cb = cb.sort_values(['BranchName', 'Date', 'Value'])

smk = SM.dropna(subset=['LST']).groupby(['BranchName', 'Code']).LST.max()   # duplicate product records: the latest count wins
g['LastCountMaster'] = [smk.get((a, c_)) for a, c_ in zip(g.BranchName, g.Code)]
g['DaysSinceCount'] = [(ASOF - v).days if v is not None and pd.notna(v) else None for v in g.LastCountMaster]
wkc = CVP[(CVP.Date >= WF) & (CVP.Date <= WT)].groupby(['BranchName', 'Code']).agg(P_=('POS', 'sum'), C_=('Cons', 'sum'))
g['RecipeNote'] = [(f'recipe consumed {wkc.loc[(a, c_)].C_:g}, recon POS {wkc.loc[(a, c_)].P_:g}' if (a, c_) in wkc.index else '') for a, c_ in zip(g.BranchName, g.Code)]
tt = t.copy()
tt['Value'] = np.where(tt.Test.str.startswith('S4'), (tt.IssuedAmended - tt.IssuedOrig) * tt.RandPU, tt.VarAmended * tt.RandPU)
bb = tt[tt.Test.str.match(r'S(7|8|11) ')].sort_values(['BranchName', 'Date'])
isf = tt[tt.Test.str.match(r'S(4a|4b|1|3) ') | tt.Test.str.startswith('S4 Issue mis')].sort_values(['BranchName', 'Test', 'Date'])
isf.loc[isf.Test.str.startswith('S4 Issue mis'), 'Value'] = 0

# ------------------------------------------------------------------ styles
F = 'Arial'
NAVY = '13233A'; HDR = PatternFill('solid', fgColor=NAVY); HF = Font(name=F, bold=True, color='FFFFFF', size=9)
WORK = PatternFill('solid', fgColor='FFF4D6'); WHF = Font(name=F, bold=True, color=NAVY, size=9); WHFILL = PatternFill('solid', fgColor='F5C96B')
BF = Font(name=F, size=9); BB = Font(name=F, size=9, bold=True); TF = Font(name=F, bold=True, size=14, color=NAVY)
H2 = Font(name=F, bold=True, size=11, color=NAVY); SF = Font(name=F, italic=True, size=9, color='5B6770')
thin = Side(style='thin', color='C9D1D3'); BOT = Border(bottom=thin)
R0 = '#,##0;(#,##0);-'; Q = '#,##0.##;(#,##0.##);-'; PC = '0%'
FAM = {'Real loss': 'F9D4CF', 'Storeroom': 'FCE0C8', 'Paperwork': 'DCEAF0', 'Receiving / invoices': 'DCE8F7', 'Reversal': 'FDEBC8',
       'Data / system': 'E6DDF6', 'Not counted': 'E7EBEC', 'Balanced': 'EEF2EE'}
STATUS = ['Open', 'Investigating', 'Escalated to Ops Exec', 'Closed - real loss, staff charged', 'Closed - real loss, key holder accountable',
          'Closed - paperwork corrected', 'Closed - invoice / GRV corrected', 'Closed - recount corrected', 'Closed - system / master data fixed', 'No action']
WORKCOLS = ['Owner (name)', 'Due date', 'Status', 'Finding - what actually happened', 'Rand charged / recovered', 'Corrected in PumpIT (Y/N)', 'Closed date', 'Head office sign-off']

wb = Workbook()
lists = wb.active; lists.title = 'Lists'
for i, s_ in enumerate(STATUS, 1): lists.cell(i, 1, s_)
for i, s_ in enumerate(['Y', 'N'], 1): lists.cell(i, 2, s_)
lists.sheet_state = 'hidden'

def header(ws, row, cols, work_from=None):
    for j, h in enumerate(cols, 1):
        c_ = ws.cell(row, j, h)
        if work_from and j >= work_from: c_.fill = WHFILL; c_.font = WHF
        else: c_.fill = HDR; c_.font = HF
        c_.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[row].height = 40

def register(name, title, sub, df, fmts=None, widths=None, famcol=None, default_status='Open', status_fn=None, tab=None):
    ws = wb.create_sheet(name)
    if tab: ws.sheet_properties.tabColor = tab
    ws['A1'] = title; ws['A1'].font = TF; ws['A2'] = sub; ws['A2'].font = SF
    cols = list(df.columns) + WORKCOLS
    wf = len(df.columns) + 1
    header(ws, 4, cols, wf)
    n = len(df)
    for i, row in enumerate(df.itertuples(index=False), 5):
        for j, v in enumerate(row, 1):
            if isinstance(v, float) and np.isnan(v): v = None
            if isinstance(v, pd.Timestamp): v = v.to_pydatetime()
            c_ = ws.cell(i, j, v); c_.font = BF; c_.border = BOT
            fm = (fmts or {}).get(df.columns[j - 1])
            if fm: c_.number_format = fm
            elif hasattr(v, 'year'): c_.number_format = 'dd-mmm'
        stv = status_fn(row) if status_fn else default_status
        for j in range(wf, wf + len(WORKCOLS)):
            c_ = ws.cell(i, j); c_.fill = WORK; c_.font = BF; c_.border = BOT
        ws.cell(i, wf + 2, stv)
        ws.cell(i, wf + 1).number_format = 'dd-mmm'; ws.cell(i, wf + 6).number_format = 'dd-mmm'; ws.cell(i, wf + 4).number_format = R0
    if famcol:
        k = list(df.columns).index(famcol) + 1
        for i in range(5, 5 + n):
            v = ws.cell(i, k).value
            if v in FAM: ws.cell(i, k).fill = PatternFill('solid', fgColor=FAM[v])
    last = 4 + max(n, 1)
    dv = DataValidation(type='list', formula1=f'=Lists!$A$1:$A${len(STATUS)}', allow_blank=True); ws.add_data_validation(dv)
    dv.add(f'{CL(wf + 2)}5:{CL(wf + 2)}{last}')
    dv2 = DataValidation(type='list', formula1='=Lists!$B$1:$B$2', allow_blank=True); ws.add_data_validation(dv2)
    dv2.add(f'{CL(wf + 5)}5:{CL(wf + 5)}{last}')
    sc = f'{CL(wf + 2)}5:{CL(wf + 2)}{last}'
    ws.conditional_formatting.add(sc, FormulaRule(formula=[f'LEFT({CL(wf + 2)}5,6)="Closed"'], fill=PatternFill('solid', fgColor='D5EDD9')))
    ws.conditional_formatting.add(sc, FormulaRule(formula=[f'LEFT({CL(wf + 2)}5,9)="Escalated"'], fill=PatternFill('solid', fgColor='F9D4CF')))
    for j, h in enumerate(cols, 1):
        ws.column_dimensions[CL(j)].width = (widths or {}).get(h, max(9, min(30, len(str(h)) + 2)))
    for h, wd in zip(WORKCOLS, [16, 10, 26, 44, 12, 11, 10, 14]): ws.column_dimensions[CL(wf + WORKCOLS.index(h))].width = wd
    ws.freeze_panes = ws.cell(5, 4)
    ws.auto_filter.ref = f'A4:{CL(len(cols))}{last}'
    return ws, wf, last

REG = {}  # sheet -> (site col, status col, value col, last row, extra)

# ------------------------------------------------------------------ Item Register (week 14-20 Sep)
g['Family'] = g.Scenario.map({m[0]: m[1] for m in MATRIX})
g['StakeUsed'] = g.S_used
ir = pd.DataFrame({'Line': g.SheetRow, 'Site': g.BranchName, 'POS system': g.Posname, 'Area': g.AreaDescription, 'Code': g.POSCode,
    'Item': g.StockItemDescription.astype(str).str.strip(), 'RandPU': g.PU, 'Recon var qty (7 days)': g.R, 'Stock take var qty': g.StakeUsed,
    'Storeroom qty (stock take - recon)': g.ST_used, 'Last week stock take qty': g.P, 'Scenario': g.Scenario, 'Family': g.Family, 'Leak bucket': g.Bucket,
    'Reason': g.Reason, 'Evidence': g.Evidence, 'Accountable': g.Accountable, 'Next step': g.NextStep, 'Priority': g.Priority,
    'Value at risk (R, selling)': g.ValueAtRisk, 'Last stock take (stock master)': g.LastCountMaster,
    f'Days since last count (at {ASOF:%d %b})': g.DaysSinceCount, 'Recipe consumption vs recon POS (week)': g.RecipeNote})
ws, wf, last = register('Item Register', f'Item register - week {WL} (the working list)',
    'Every recon item on the Stock Recon Detail for the week, diagnosed. Work the yellow columns. Values at the recon RandPU (selling).',
    ir, {'RandPU': '#,##0.00', 'Recon var qty (7 days)': Q, 'Stock take var qty': Q, 'Storeroom qty (stock take - recon)': Q, 'Last week stock take qty': Q, 'Value at risk (R, selling)': R0},
    {'Site': 22, 'Area': 16, 'Item': 32, 'Reason': 70, 'Evidence': 50, 'Accountable': 24, 'Next step': 45, 'Leak bucket': 22, 'Family': 14, 'Last stock take (stock master)': 12, 'Recipe consumption vs recon POS (week)': 30},
    famcol='Family', status_fn=lambda row: 'No action' if row[11] == 'OK' else 'Open', tab='F5C96B')
REG['Item Register'] = dict(site='B', status=CL(wf + 2), value='T', bucket='N', scen='L', pri='S', last=last)

# ------------------------------------------------------------------ Opening vs Closing
oc = pd.DataFrame({'Site': cb.BranchName, 'Area': cb.AreaDescription, 'Code': cb.POSCode, 'Item': cb.StockItemDescription, 'Date': cb.Date, 'Shift': cb.ShiftNo,
    'Previous closing (staff count)': cb.PrevCloseStaff, 'Opening captured': cb.QtyOpen_Original, 'Difference': cb.Break, 'Direction': cb.Direction,
    'Previous closing after balancing': cb.PrevCloseBal, 'Opening after balancing': cb.QtyOpen_Amended, 'How balancing handled it': cb.Repair,
    'RandPU': cb.RandPU, 'Value (R, selling)': cb.Value})
ws, wf, last = register('Opening vs Closing', f'Every shift where the opening did not equal the previous closing - {PL}',
    'Rule: amend the OPENING to the previous closing; the closing count stands. "WRONG" = balancing changed the closing to match the opening. Explain every line.',
    oc, {'Difference': Q, 'Previous closing (staff count)': Q, 'Opening captured': Q, 'Previous closing after balancing': Q, 'Opening after balancing': Q, 'RandPU': '#,##0.00', 'Value (R, selling)': R0},
    {'Site': 22, 'Item': 32, 'Direction': 28, 'How balancing handled it': 30}, tab='C2410C')
REG['Opening vs Closing'] = dict(site='A', status=CL(wf + 2), value='O', dirc='J', rep='M', last=last)

# ------------------------------------------------------------------ Balancing breaches
bz = pd.DataFrame({'Site': bb.BranchName, 'Area': bb.Area, 'Code': bb.Code, 'Item': bb.Item, 'Date': bb.Date, 'Shift': bb.Shift, 'Test': bb.Test,
    'Variance - staff count': bb.VarOrig, 'Variance - after balancing': bb.VarAmended, 'Value (R, selling)': bb.Value, 'Detail': bb.Detail})
ws, wf, last = register('Balancing Breaches', f'Balancing breaches - {PL}',
    'S7 over turned into a short; S8 short increased; S11 latest closing count changed. Each breaks the balancing rule and must be reversed and explained.',
    bz, {'Variance - staff count': Q, 'Variance - after balancing': Q, 'Value (R, selling)': R0}, {'Site': 22, 'Item': 32, 'Test': 32, 'Detail': 60}, tab='B3261E')
REG['Balancing Breaches'] = dict(site='A', status=CL(wf + 2), value='J', test='G', last=last)

# ------------------------------------------------------------------ Issue sheet failures
iz = pd.DataFrame({'Site': isf.BranchName, 'Area': isf.Area, 'Code': isf.Code, 'Item': isf.Item, 'Date': isf.Date, 'Shift': isf.Shift, 'Test': isf.Test,
    'Issued - captured on shift': isf.IssuedOrig, 'Issued - after balancing': isf.IssuedAmended, 'Pack size': isf.Multiple, 'Value (R, selling)': isf.Value, 'Detail': isf.Detail})
ws, wf, last = register('Issue Sheet Failures', f'Issue-sheet failures - {PL}',
    'S4a issue added at balancing (not written up on the shift); S4b issue reduced (over-captured / captured twice); S4 mis-keyed; S1 over = a whole pack still standing; S3 short = a whole pack.',
    iz, {'Issued - captured on shift': Q, 'Issued - after balancing': Q, 'Value (R, selling)': R0}, {'Site': 22, 'Item': 32, 'Test': 34, 'Detail': 60}, tab='2B6CB0')
REG['Issue Sheet Failures'] = dict(site='A', status=CL(wf + 2), value='K', test='G', last=last)

# ------------------------------------------------------------------ Control actions (management)
nwk = g.groupby('BranchName').agg(lines=('R', 'size'), nc=('Scenario', lambda s_: (s_ == 'NC').sum()))
oc_s = cb[cb.Break < 0].groupby('BranchName').agg(n=('Value', 'size'), v=('Value', 'sum'), wrong=('Repair', lambda s_: (s_ != 'Correct - opening amended').sum()))
br_s = bb[bb.Test.str.match(r'S(7|8) ')].groupby('BranchName').agg(n=('Value', 'size'), v=('Value', 'sum'))
la_s = isf[isf.Test.str.startswith('S4a')].groupby('BranchName').agg(n=('Value', 'size'), v=('Value', 'sum'))
CA = []
def add(pri, site, area, action, why, role, due):
    CA.append(dict(Priority=pri, Site=site, Area=area, Action=action, Why=why, Role=role, Due=due))
add('1', 'All sites', 'Recon control', 'Every opening must equal the previous closing. Where it does not, amend the OPENING; never amend a closing to match an opening.',
    f'{int(oc_s.n.sum())} shifts opened below the previous closing {PLs}; R{-oc_s.v.sum():,.0f} vanished between shifts; {int(oc_s.wrong.sum())} fixed the wrong way or still open.', 'Ops Exec + site managers', 'Immediately')
for s_, row in oc_s.sort_values('v').iterrows():
    if row.n >= 10:
        add('1', s_, 'Recon control', f'Explain every opening below the previous closing on the Opening vs Closing tab ({int(row.n)} shifts).',
            f'R{-row.v:,.0f} vanished between shifts; {int(row.wrong)} of {int(row.n)} fixed by changing the closing or left open.', 'Site manager', 'Within 48 hours')
for s_, row in br_s.sort_values('v').iterrows():
    if row.n >= 10:
        add('1', s_, 'Balancing', f'Reverse and explain {int(row.n)} balancing breaches (over made short / short increased).', f'R{-row.v:,.0f} at selling price moved against the rule.', 'Site manager', 'Within 48 hours')
for s_, row in la_s.sort_values('v', ascending=False).iterrows():
    if row.n >= 15:
        add('2', s_, 'Issue sheets', f'Issues are to be written up on the shift. {int(row.n)} issues were only captured at balancing.', f'R{row.v:,.0f} at selling price issued without paperwork at the time.', 'Site manager', 'This week')
for _, x_ in irp.iterrows():
    if 'breached' in x_.Verdict and x_.Code == 'V7':
        add('1', x_.BranchName, 'Steers BOH', 'Issue every potato delivery in full on the day received (issue = GRV). Key holder to account for the shortfall.',
            f'{x_.Compliance}; received {x_.GRV:,.0f} kg, issued {x_.Issued:,.0f} kg; site count short {-x_.SiteVar:,.1f} kg (R{-x_.SiteVarR:,.0f} at cost) {PLs}.', 'Site manager / key holder', 'Immediately')
for _, x_ in irp.iterrows():
    if 'breached' in x_.Verdict and x_.Code in ('B1', 'SBBB'):
        add('2', x_.BranchName, 'Steers BOH', f'Issue buns ({x_.Code}) in full on delivery - stop holding them in a store.', f'{x_.Compliance} {PLs}.', 'Site manager', 'This week')
for s_, row in nwk.iterrows():
    if row.nc / row.lines >= 0.2:
        add('2', s_, 'Counting', f'Count every recon item on Monday. {int(row.nc)} of {int(row.lines)} recon lines had no stock take in {WLs}.',
            'A recon item not counted cannot be split between the counter and the storeroom.', 'Stock controller / site manager', 'Monday')
_smk = SM.dropna(subset=['LST']).groupby(['BranchName', 'Code']).LST.max()
_wb = g[g.Posname == 'WINBRANCH'][['BranchName', 'Code', 'AreaDescription', 'StockItemDescription']].drop_duplicates(['BranchName', 'Code'])
_wb['LST'] = [_smk.get((a_, c_)) for a_, c_ in zip(_wb.BranchName, _wb.Code)]
_wb['Days'] = [(ASOF - v).days if v is not None and pd.notna(v) else None for v in _wb.LST]
for s_, x_ in _wb[pd.to_numeric(_wb.Days) > 30].groupby('BranchName'):
    if len(x_) >= 5:
        _ar = ', '.join(x_.AreaDescription.value_counts().head(3).index)
        add('1', s_, _ar, f'Count the {len(x_)} recon items not counted for more than 30 days - now, not next Monday.',
            f'Oldest last count {x_.LST.min():%d %B %Y} (stock master). An item that is not counted cannot be split between the counter and the storeroom.', 'Site manager / stock controller', 'Immediately')
if len(CVP):
    _cv = CVP.groupby(['BranchName', 'Code']).agg(Item=('Item', 'first'), Area=('Area', 'first'), POS=('POS', 'sum'), Cons=('Cons', 'sum')).reset_index()
    _cv['Gap'] = _cv.POS - _cv.Cons
    _cv = _cv[_cv.Gap.abs() > np.maximum(20, 0.2 * _cv.Cons.abs())].reindex(_cv.Gap.abs().sort_values(ascending=False).index).dropna(subset=['Gap']).head(6)
    for _, x_ in _cv.iterrows():
        add('2', x_.BranchName, x_.Area, f'Fix the recon / recipe link for {str(x_.Item).strip()} ({x_.Code}): recipes consumed {x_.Cons:,.0f} over {PLs}, the recon POS saw {x_.POS:,.0f}.',
            'The recon shows a false short at the counter.' if x_.Gap < 0 else 'A selling code is linked to the recon but not to the recipe - the stock book stays too high and the stock take shows a false short.',
            'IT / master data', 'This week')
for s_, x_ in g[g.Scenario == 'E1'].groupby('BranchName'):
    add('2', s_, ', '.join(sorted(set(x_.AreaDescription)))[:40], f'Link the selling codes for {", ".join(x_.StockItemDescription.astype(str).str.strip().head(5))} to their recon items.',
        f'{len(x_)} line(s) this week show stock moving with no till sales linked - a system fault, not theft.', 'IT / master data', 'This week')
for s_, x_ in g[g.Scenario == 'E2'].groupby('BranchName'):
    add('2', s_, ', '.join(sorted(set(x_.AreaDescription)))[:40], f'Recount before posting: {", ".join(x_.StockItemDescription.astype(str).str.strip().head(6))}.',
        'Stock take variance more than 10 times throughput - unit of measure, code or recipe fault.', 'Stock controller', 'This week')
for s_, x_ in g[(g.Posname == 'NAMOS') & (g.Scenario == 'NC')].groupby('BranchName'):
    add('3', s_, 'Shop (NAMOS)', 'Link a NAMOS stock take to the shop recon.', f'{len(x_)} shop recon lines cannot be reconciled.', 'IT', '30 days')
_am = int(g.Evidence.str.contains('Aura stk 01').sum())
if _am:
    add('3', 'All QSR sites', 'Aura', f'Correct the Aura stock take feed on the weekly sheet - it disagrees with the Aura count register on {_am} lines.', 'The register is used in this book; the sheet misstates QSR storeroom results.', 'IT / data', '30 days')
add('3', 'All WinBranch sites', 'Data', 'Include zero-variance count lines in the WinBranch stock take export.', 'Without them a counted-and-balanced item looks "not counted".', 'IT / data', '30 days')
add('3', 'All sites', 'Recon master', 'Remove virtual lines (airtime, electricity, lotto, refill water) from the recon item lists.', f'{int((g.Scenario == "V").sum())} virtual lines on the recon this week.', 'IT / master data', '30 days')
add('3', 'All sites', 'Recon master', 'Load the issue multiples; reject issues that are not whole packs.', f'{int((t.Test.str.startswith("S5")).sum())} issues {PLs} were not in the pack multiple.', 'IT / master data', '30 days')
add('3', 'All sites', 'ERP', 'Lock the opening to the previous closing; handover difference signed by both staff; balancing may never change a closing to match an opening.', 'Removes the manipulation route.', 'IT / ERP', 'ERP build')
add('3', 'All sites', 'ERP', 'Put the cashier / staff ID on every recon line.', 'A short can then be charged to a person, not a shift number.', 'IT / ERP', 'ERP build')
add('3', 'All sites', 'ERP', 'A stock take over or a reversal of last week cannot be posted without a reason code.', f'{int((g.Scenario == "RV").sum())} reversals and {int((g.Scenario == "D2").sum())} storeroom overs this week.', 'IT / ERP', 'ERP build')
ca = pd.DataFrame(CA); ca.insert(0, 'No.', range(1, len(ca) + 1))
ca.columns = ['No.', 'Priority (1 = now)', 'Site', 'Area', 'Action required', 'Why - the evidence', 'Responsible role', 'Due']
ws, wf, last = register('Control Actions', 'Control actions - management', 'Site-level and system actions that stop the leaks at source. Priority 1 = do now; 2 = this week; 3 = within 30 days / ERP.',
    ca, None, {'Action required': 60, 'Why - the evidence': 60, 'Site': 22, 'Responsible role': 22, 'Area': 16}, tab='13233A')
for i in range(5, last + 1):
    for j in (5, 6): wb['Control Actions'].cell(i, j).alignment = Alignment(wrap_text=True, vertical='top')
REG['Control Actions'] = dict(site='C', status=CL(wf + 2), last=last)

# ------------------------------------------------------------------ Potatoes & buns
pb = irp[['BranchName', 'Item', 'Code', 'GRV', 'Issued', 'GRV not issued', 'POS', 'SiteVar', 'SiteVarR', 'Compliance', 'Verdict']].copy()
pb.columns = ['Site', 'Item', 'Code', 'Received (GRV)', 'Issued to floor', 'Received not issued', 'POS (recipe)', 'Site count var qty', 'Site count var (R, cost)', 'Deliveries issued as received', 'Verdict']
ws, wf, last = register('Potatoes & Buns', f'Steers potatoes and buns - issued on receipt? {PL}',
    'Potatoes (V7) and buns (B1 80g, SBBB 50g) must be issued in full on the day received; storeroom stock is nil by design. Site count per the Aura count register.',
    pb, {'Received (GRV)': Q, 'Issued to floor': Q, 'Received not issued': Q, 'POS (recipe)': Q, 'Site count var qty': Q, 'Site count var (R, cost)': R0},
    {'Site': 22, 'Verdict': 60, 'Deliveries issued as received': 16}, tab='C2410C')
r0 = last + 3
d_ = ird.copy(); d_.columns = ['Site', 'Area', 'Code', 'Item', 'Date', 'Received (GRV)', 'Invoice', 'Issued', 'Result', 'GRV exception']
ws.cell(r0 - 1, 1, 'Delivery-by-delivery detail').font = H2
header(ws, r0, list(d_.columns))
for i, row in enumerate(d_.itertuples(index=False), r0 + 1):
    for j, v in enumerate(row, 1):
        if isinstance(v, pd.Timestamp): v = v.to_pydatetime()
        c_ = ws.cell(i, j, v); c_.font = BF
        if hasattr(v, 'year'): c_.number_format = 'dd-mmm'

# ------------------------------------------------------------------ Count compliance (week 14-20, corrected)
cmp_ = g.assign(Counted=g.Scenario.ne('NC')).groupby(['BranchName', 'Posname']).agg(Lines=('R', 'size'), Counted=('Counted', 'sum')).reset_index()
cmp_['Not counted'] = cmp_.Lines - cmp_.Counted
nc = g[g.Scenario == 'NC'][['BranchName', 'Posname', 'AreaDescription', 'POSCode', 'StockItemDescription', 'R', 'PU', 'Reason']].copy()
nc.columns = ['Site', 'POS system', 'Area', 'Code', 'Item', 'Recon var qty (7 days)', 'RandPU', 'Note']
ws, wf, last = register('Count Compliance', f'Recon items not counted - week {WL}',
    'Counted = a stock take in the week (WinBranch stock take detail, or a last stock take date inside the week; Aura stk 01 and Pilot registers).',
    nc, {'Recon var qty (7 days)': Q, 'RandPU': '#,##0.00'}, {'Site': 22, 'Item': 32, 'Note': 70}, tab='7A868C')
REG['Count Compliance'] = dict(site='A', status=CL(wf + 2), last=last)

# ------------------------------------------------------------------ Period items (reference)
pi = p.copy()
pi['Family'] = pi.Class.map({m[0]: m[1] for m in MATRIX})
fl = w.groupby(['BranchName', 'Code']).Flags.apply(lambda s_: ' | '.join(sorted({f for q in s_ for f in q.split(' | ') if f}))).rename('Evidence')
pi = pi.merge(fl, on=['BranchName', 'Code'], how='left')
pz = pd.DataFrame({'Site': pi.BranchName, 'Area': pi.Area, 'Code': pi.Code, 'Item': pi.Item, 'Site counts': pi.Counts, 'POS qty': pi.POS, 'Issued qty': pi.Issued,
    'Recon var qty - staff count': pi.R_orig, 'Recon var qty - balanced': pi.R, 'Site count var qty': pi.S, 'Unit cost': pi.UnitCost,
    'Recon var R - staff count': pi.R_orig_value, 'Recon var R - balanced': pi.R_value, 'Site count var R': pi.S_value, 'Storeroom R': pi.Storeroom_value,
    'Scenario': pi.Class, 'Family': pi.Family, 'Scenario name': pi.Scenario, 'Evidence': pi.Evidence})
pz = pz.sort_values(['Site', 'Scenario'])
ws, wf, last = register(PTAB, f'Item diagnosis over {PL}, at cost',
    'Count-to-count windows netted per item. Only items with a non-zero stock take variance appear (export limitation). Reference for the weekly register.',
    pz, {k: Q for k in ['POS qty', 'Issued qty', 'Recon var qty - staff count', 'Recon var qty - balanced', 'Site count var qty']} | {k: R0 for k in ['Recon var R - staff count', 'Recon var R - balanced', 'Site count var R', 'Storeroom R']} | {'Unit cost': '#,##0.00'},
    {'Site': 22, 'Item': 32, 'Scenario name': 30, 'Evidence': 60, 'Family': 14}, famcol='Family', status_fn=lambda row: 'No action' if row[15] == 'OK' else 'Open', tab='A7B3B8')

# ------------------------------------------------------------------ Cross-site
xs = x[['Code', 'Item', 'Area', 'Sites', 'ShortSites', 'OverSites', 'MedianPct', 'Direction', 'Likely', 'SiteList']].copy()
xs.columns = ['Code', 'Item', 'Area', 'Sites selling', 'Sites short >3%', 'Sites over >3%', 'Median variance % of POS', 'Direction', 'Likely cause', 'Sites']
ws, wf, last = register('Cross-Site Patterns', f'Same item, same direction at 3 or more sites - {PL}', 'A pattern across most sites points to a recipe card, pack size or PLU-link fault, not theft at every site.',
    xs, {'Median variance % of POS': PC}, {'Item': 32, 'Likely cause': 50, 'Sites': 60}, tab='6D3FC0')

# ------------------------------------------------------------------ Count recency (stock master last stock take)
BANDS = ['0-2 days', '3-9 days', '10-16 days', '17-30 days', '31-90 days', 'over 90 days', 'no date on master']
rc = g[['BranchName', 'Posname', 'AreaDescription', 'POSCode', 'StockItemDescription', 'LastCountMaster', 'DaysSinceCount']].copy()
rc = rc[rc.Posname == 'WINBRANCH']
_b = pd.cut(pd.to_numeric(rc.DaysSinceCount), [-1, 2, 9, 16, 30, 90, 1e6], labels=BANDS[:-1]).astype(object)
rc['Last count'] = _b.where(_b.notna(), 'no date on master')
rc['DaysSinceCount'] = pd.to_numeric(rc.DaysSinceCount)
old_ = rc[(rc.DaysSinceCount > 16) | rc.DaysSinceCount.isna()].sort_values(['BranchName', 'DaysSinceCount'], ascending=[True, False])
oz = old_[['BranchName', 'AreaDescription', 'POSCode', 'StockItemDescription', 'LastCountMaster', 'DaysSinceCount', 'Last count']]
oz.columns = ['Site', 'Area', 'Code', 'Item', 'Last stock take', 'Days since last count', 'Band']
ws, wf, last = register('Count Recency', f'When was each recon item last counted? (stock master, as at {ASOF:%d %B %Y})',
    'Last stock take date per recon item from the ALL SITES STOCK master (WinBranch sites). It shows only the most recent count. Listed below: items not counted for more than 16 days, or with no date.',
    oz, {'Days since last count': '#,##0'}, {'Site': 22, 'Item': 34, 'Area': 22, 'Band': 18}, tab='7A868C')
REG['Count Recency'] = dict(site='A', status=CL(wf + 2), last=last)
bands = BANDS
ct = pd.crosstab(rc.BranchName, rc['Last count']).reindex(columns=bands, fill_value=0)
r0 = last + 3; ws.cell(r0 - 1, 1, 'Recon items by last count - WinBranch sites').font = H2
header(ws, r0, ['Site'] + bands + ['Recon items'])
for i, (s_, row_) in enumerate(ct.iterrows(), r0 + 1):
    ws.cell(i, 1, s_).font = BF
    for j, v in enumerate(row_.values, 2): ws.cell(i, j, int(v)).font = BF
    ws.cell(i, len(bands) + 2, f'=SUM(B{i}:{CL(len(bands) + 1)}{i})').font = BB
_nosm = sorted(set(rc.BranchName) - set(SM.BranchName))
ws.cell(r0 + len(ct) + 2, 1, ('Not on the stock master: ' + ', '.join(_nosm) + '. ' if _nosm else '') + 'QSR (Aura) and Baobab Inn (Pilot) items are counted in their own registers; NAMOS shop lines have no stock take link.').font = SF

# ------------------------------------------------------------------ Recipe consumption vs recon POS
cv = CVP.copy()
ca_ = cv.groupby(['BranchName', 'Code']).agg(Item=('Item', 'first'), Area=('Area', 'first'), Days=('Date', 'nunique'), POS=('POS', 'sum'), Cons=('Cons', 'sum'),
    NoPOS=('st', lambda s_: (s_ == 'recipe used, recon POS nil').sum()), Differ=('st', lambda s_: (s_ == 'both, differ >10%').sum())).reset_index()
ca_['Gap'] = ca_.POS - ca_.Cons
def meaning(x_):
    if abs(x_.Gap) <= max(2, 0.05 * x_.Cons) and x_.NoPOS == 0: return 'Recipe and recon agree'
    if x_.Gap < 0: return 'Recipe reduces the stock book by more than the recon POS expects: the recon shows a false short at the counter, or the recon link misses some selling codes'
    return 'Recon POS expects more than the recipe takes off the stock book: a selling code is linked to the recon but not to the recipe - the stock book stays too high and the stock take shows a false short'
ca_['What it means'] = ca_.apply(meaning, axis=1)
ca_ = ca_.sort_values(['What it means', 'Gap'])
cz = ca_[['BranchName', 'Area', 'Code', 'Item', 'Days', 'Cons', 'POS', 'Gap', 'NoPOS', 'Differ', 'What it means']]
cz.columns = ['Site', 'Area', 'Code', 'Item', 'Days', 'Recipe consumption (WinBranch raw material / Pilot theoretical)', 'Recon POS quantity', 'Gap (recon - recipe)', 'Days recipe used, recon POS nil', 'Days differ >10%', 'What it means']
ws, wf, last = register('Recipe vs Recon', f'Recipe consumption vs recon POS - recon items used in recipes, {PL}',
    'Raw material consumption (recipe explosion of every sale) against the POS quantity the recon uses. They should agree every day. A gap is a link or recipe fault, not a loss.',
    cz, {'Recipe consumption (WinBranch raw material / Pilot theoretical)': Q, 'Recon POS quantity': Q, 'Gap (recon - recipe)': Q}, {'Site': 22, 'Item': 32, 'What it means': 80}, tab='6D3FC0')
REG['Recipe vs Recon'] = dict(site='A', status=CL(wf + 2), last=last)

# ------------------------------------------------------------------ Scenario matrix
sm = wb.create_sheet('Scenario Matrix', 0); sm.sheet_properties.tabColor = '0E6E6A'
sm['A1'] = 'Scenario matrix - how every variance is read'; sm['A1'].font = TF
sm['A2'] = 'Recon point = stock in the staff member\'s hands (opening + issues - closing vs POS). Site stock take = the whole site, storeroom included. Nil = within 1 unit (0.2 for weighed) or 2% of POS. Same size = within 15%.'; sm['A2'].font = SF
mh = ['Code', 'Family', 'Scenario', 'Recon point', 'Site stock take', 'Other test', 'What really happened (most likely first)', 'Evidence to pull', 'Accountable', 'Charge staff?', 'Action to close', 'ERP control that prevents it']
header(sm, 4, mh)
for i, m in enumerate(MATRIX, 5):
    for j, v in enumerate(m, 1):
        c_ = sm.cell(i, j, v); c_.font = BB if j <= 3 else BF; c_.alignment = Alignment(wrap_text=True, vertical='top'); c_.border = BOT
    if m[1] in FAM:
        for j in (1, 2): sm.cell(i, j).fill = PatternFill('solid', fgColor=FAM[m[1]])
for col, wd in zip('ABCDEFGHIJKL', [6, 14, 26, 16, 16, 16, 50, 36, 22, 16, 34, 40]): sm.column_dimensions[col].width = wd
r1 = 6 + len(MATRIX)
sm.cell(r1, 1, 'Shift-level tests (balancing screen)').font = H2
header(sm, r1 + 1, ['Test', 'What it is', 'What it means', 'Action', 'Accountable', 'ERP control'])
for i, s_ in enumerate(SHIFT_TESTS, r1 + 2):
    for j, v in enumerate(s_, 1):
        c_ = sm.cell(i, j, v); c_.font = BB if j <= 2 else BF; c_.alignment = Alignment(wrap_text=True, vertical='top'); c_.border = BOT
r2 = r1 + 3 + len(SHIFT_TESTS)
sm.cell(r2, 1, 'Rules every investigation follows').font = H2
for i, rl in enumerate(RULES, r2 + 1):
    sm.cell(i, 1, i - r2).font = BF; c_ = sm.cell(i, 2, rl); c_.font = BF
    sm.merge_cells(start_row=i, start_column=2, end_row=i, end_column=9); c_.alignment = Alignment(wrap_text=True, vertical='top'); sm.row_dimensions[i].height = 26
sm.freeze_panes = 'D5'

# ------------------------------------------------------------------ Site scorecard (live formulas)
scd = wb.create_sheet('Site Scorecard', 0); scd.sheet_properties.tabColor = '0E6E6A'
scd['A1'] = 'Site scorecard - live'; scd['A1'].font = TF
scd['A2'] = 'Counts and values update as the yellow Status columns on the registers are worked. Open = Open, Investigating or Escalated.'; scd['A2'].font = SF
IRr = REG['Item Register']; OCr = REG['Opening vs Closing']; BRr = REG['Balancing Breaches']; ISr = REG['Issue Sheet Failures']
rng = lambda sheet, col, lst: f"'{sheet}'!${col}$5:${col}${lst}"
def openc(sheet, reg, site_ref, extra=''):
    s_, st_ = rng(sheet, reg['site'], reg['last']), rng(sheet, reg['status'], reg['last'])
    return '+'.join(f'COUNTIFS({s_},{site_ref},{st_},"{k}"{extra})' for k in ['Open', 'Investigating', 'Escalated to Ops Exec'])
cols = ['Site', f'Week {WLs}: lines to investigate', 'Week: open lines', 'Week: closed lines', 'Week: % closed', 'Week: value at risk (R)', 'Week: storeroom loss (R)',
        'Week: loss at recon point (R)', 'Week: not counted lines', f'Opening below previous closing (shifts, {PLs})', 'Stock vanished between shifts (R)',
        'Opening/closing lines still open', f'Balancing breaches ({PLs})', 'Balancing breaches still open', f'Issue-sheet failures ({PLs})', 'Issue-sheet failures still open',
        'Recon items last counted over 16 days ago (stock master)']
header(scd, 4, cols)
for i, s_ in enumerate(SITES, 5):
    a = f'$A{i}'
    scd.cell(i, 1, s_)
    S_I, ST_I, V_I, B_I = rng('Item Register', IRr['site'], IRr['last']), rng('Item Register', IRr['status'], IRr['last']), rng('Item Register', IRr['value'], IRr['last']), rng('Item Register', IRr['bucket'], IRr['last'])
    scd.cell(i, 2, f'=COUNTIFS({S_I},{a},{B_I},"<>Balanced")')
    scd.cell(i, 3, '=' + openc('Item Register', IRr, a))
    scd.cell(i, 4, f'=COUNTIFS({S_I},{a},{ST_I},"Closed*")')
    scd.cell(i, 5, f'=IF(B{i}=0,0,D{i}/B{i})').number_format = PC
    scd.cell(i, 6, f'=SUMIFS({V_I},{S_I},{a},{B_I},"<>Balanced",{B_I},"<>Virtual line")').number_format = R0
    scd.cell(i, 7, f'=SUMIFS({V_I},{S_I},{a},{B_I},"Storeroom loss")').number_format = R0
    scd.cell(i, 8, f'=SUMIFS({V_I},{S_I},{a},{B_I},"Loss at recon point")').number_format = R0
    scd.cell(i, 9, f'=COUNTIFS({S_I},{a},{B_I},"Not counted")')
    S_O, D_O, V_O = rng('Opening vs Closing', 'A', OCr['last']), rng('Opening vs Closing', OCr['dirc'], OCr['last']), rng('Opening vs Closing', OCr['value'], OCr['last'])
    scd.cell(i, 10, f'=COUNTIFS({S_O},{a},{D_O},"Opening BELOW*")')
    scd.cell(i, 11, f'=-SUMIFS({V_O},{S_O},{a},{D_O},"Opening BELOW*")').number_format = R0
    scd.cell(i, 12, '=' + openc('Opening vs Closing', OCr, a))
    scd.cell(i, 13, f'=COUNTIF({rng("Balancing Breaches", "A", BRr["last"])},{a})')
    scd.cell(i, 14, '=' + openc('Balancing Breaches', BRr, a))
    scd.cell(i, 15, f'=COUNTIF({rng("Issue Sheet Failures", "A", ISr["last"])},{a})')
    scd.cell(i, 16, '=' + openc('Issue Sheet Failures', ISr, a))
    scd.cell(i, 17, f"=COUNTIFS('Count Recency'!$A$5:$A${REG['Count Recency']['last']},{a},'Count Recency'!$G$5:$G${REG['Count Recency']['last']},\"<>no date on master\")")
    for j in range(1, 18): scd.cell(i, j).font = BF; scd.cell(i, j).border = BOT
tr = 5 + len(SITES); scd.cell(tr, 1, 'TOTAL').font = BB
for j in range(2, 18):
    L_ = CL(j)
    c_ = scd.cell(tr, j, f'=SUM({L_}5:{L_}{tr - 1})' if j != 5 else f'=IF(B{tr}=0,0,D{tr}/B{tr})'); c_.font = BB
    c_.number_format = PC if j == 5 else (R0 if j in (6, 7, 8, 11) else '#,##0')
scd.column_dimensions['A'].width = 28
for j in range(2, 18): scd.column_dimensions[CL(j)].width = 13
scd.row_dimensions[4].height = 64
# bucket table by site
b0 = tr + 3
scd.cell(b0 - 1, 1, f'Week {WLs}: value at risk by leak bucket (R, selling)').font = H2
bks = [b for b in BUCKET_ORDER if b not in ('Balanced', 'Production item')]
header(scd, b0, ['Site'] + bks)
for i, s_ in enumerate(SITES, b0 + 1):
    scd.cell(i, 1, s_).font = BF
    for j, bk in enumerate(bks, 2):
        c_ = scd.cell(i, j, f'=SUMIFS({rng("Item Register", IRr["value"], IRr["last"])},{rng("Item Register", IRr["site"], IRr["last"])},$A{i},{rng("Item Register", IRr["bucket"], IRr["last"])},"{bk}")')
        c_.number_format = R0; c_.font = BF
t2 = b0 + 1 + len(SITES); scd.cell(t2, 1, 'TOTAL').font = BB
for j in range(2, 2 + len(bks)):
    c_ = scd.cell(t2, j, f'=SUM({CL(j)}{b0 + 1}:{CL(j)}{t2 - 1})'); c_.number_format = R0; c_.font = BB
scd.row_dimensions[b0].height = 44
scd.freeze_panes = 'B5'

exec(open(os.path.join(HERE, 'findings.py')).read())
order = ['Read Me', 'Findings', 'Site Scorecard', 'Scenario Matrix', 'Control Actions', 'Item Register', 'Opening vs Closing', 'Balancing Breaches',
         'Issue Sheet Failures', 'Potatoes & Buns', 'Count Compliance', 'Count Recency', 'Recipe vs Recon', PTAB, 'Cross-Site Patterns', 'Lists']
wb._sheets = [wb[n] for n in order]
wb.active = 0
wb.save(OUT); print('saved', OUT)
